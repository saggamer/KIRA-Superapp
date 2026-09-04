#!/usr/bin/env python3
"""Non-destructive host-level release probes for KIRA OS.

This script deliberately avoids mutations outside KIRA's own stress workspace.
It is intended to be run on the real macOS host so MLX, networking, and native
frameworks are exercised instead of mocked.
"""

from __future__ import annotations

import base64
import json
import os
import shutil
import sys
import time
import traceback
from pathlib import Path


APP_DIR = Path(__file__).resolve().parents[1]
REPORT_DIR = APP_DIR / "stress_lab" / "ultra_release"
REPORT_PATH = REPORT_DIR / "host_probe.json"
PROBE_WORKSPACE = REPORT_DIR / "workspace"
sys.path.insert(0, str(APP_DIR))


def timed(name, operation):
    started = time.perf_counter()
    try:
        detail = operation()
        ok = detail.get("ok", True) if isinstance(detail, dict) else bool(detail)
        return {
            "name": name,
            "ok": bool(ok),
            "status": "passed" if ok else "failed",
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "detail": detail,
        }
    except Exception as error:
        message = str(error)
        environment_block = any(marker in message.lower() for marker in (
            "no metal device available",
            "network is unreachable",
            "operation not permitted",
        ))
        return {
            "name": name,
            "ok": False,
            "status": "blocked" if environment_block else "failed",
            "elapsed_seconds": round(time.perf_counter() - started, 3),
            "error": message,
            "trace": traceback.format_exc(),
        }


def artifact_probe(brain):
    from PIL import Image
    from docx import Document
    from pptx import Presentation

    image_path = PROBE_WORKSPACE / "release_probe.png"
    Image.new("RGB", (960, 540), (27, 94, 132)).save(image_path)
    pptx_block = f"""TITLE: KIRA Ultra Release Probe
SUBTITLE: Host-verified presentation output
THEME: dark cyan
SLIDE: Verification Scope
BODY: This deck verifies native PowerPoint generation, substantive slide content, embedded local imagery, and direct source hyperlinks.
IMAGE: {image_path}
LINK: Python documentation | https://docs.python.org/3/
SLIDE: Evidence
BODY: The artifact exists on disk, opens as a valid Office package, and contains an image plus a clickable external relationship.
LINK: KIRA OS repository | https://github.com/saggamer/KIRA-OS
SLIDE: Release Gate
BODY: Generation is accepted only after the resulting package can be parsed and its expected content is present.
"""
    pptx_result = brain._generate_native_pptx(pptx_block)
    pptx_path = brain._extract_artifact_path_from_text(pptx_result)
    presentation = Presentation(pptx_path) if pptx_path else None
    image_count = 0
    hyperlink_count = 0
    if presentation:
        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) == 13:
                    image_count += 1
                if not getattr(shape, "has_text_frame", False):
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    hyperlink_count += sum(1 for run in paragraph.runs if run.hyperlink.address)

    docx_block = """TITLE: KIRA Ultra Release Probe
SUBTITLE: Host-verified Word output
CONTENT:
# Verification
KIRA OS generated this document through its native Word artifact branch as a
bounded host-level release check. The check confirms that the output is a real
Office Open XML document, not a textual claim that a document was created.

## Package integrity
The resulting file must exist on disk and open through the standard Python DOCX
reader. It must preserve a meaningful title, multiple substantive paragraphs,
and enough readable content to be useful outside the chat interface.

## Evidence contract
The document includes an external source link and records the acceptance rules
used by the probe. A file is rejected when generation returns an error, when the
package cannot be parsed, or when the artifact quality report finds thin or
placeholder content.

## Release result
Passing this check demonstrates that KIRA OS can create and verify a new Word
artifact inside its isolated probe workspace. It does not infer success from a
filename alone; both generation and post-write inspection must agree.
LINK: Python documentation | https://docs.python.org/3/
"""
    docx_result = brain._generate_native_docx(docx_block)
    docx_path = brain._extract_artifact_path_from_text(docx_result)
    document = Document(docx_path) if docx_path else None
    docx_text = "\n".join(p.text for p in document.paragraphs) if document else ""

    pptx_ok = bool(
        pptx_path
        and Path(pptx_path).exists()
        and presentation
        and len(presentation.slides) >= 4
        and image_count >= 1
        and hyperlink_count >= 2
    )
    docx_quality = (
        brain._docx_quality_report(docx_path, docx_block)
        if docx_path and Path(docx_path).exists()
        else {"ok": False, "issues": ["No DOCX artifact was created."]}
    )
    docx_ok = bool(
        docx_path
        and Path(docx_path).exists()
        and document
        and "DOCX generated:" in docx_result
        and docx_quality.get("ok", False)
        and "Verification" in docx_text
    )
    return {
        "ok": pptx_ok and docx_ok,
        "pptx": {
            "ok": pptx_ok,
            "path": pptx_path,
            "slides": len(presentation.slides) if presentation else 0,
            "images": image_count,
            "hyperlinks": hyperlink_count,
            "generator_result": pptx_result,
        },
        "docx": {
            "ok": docx_ok,
            "path": docx_path,
            "paragraphs": len(document.paragraphs) if document else 0,
            "characters": len(docx_text),
            "quality": docx_quality,
            "generator_result": docx_result,
        },
    }


def main():
    from interface import KiraBrain

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if PROBE_WORKSPACE.exists():
        shutil.rmtree(PROBE_WORKSPACE)
    PROBE_WORKSPACE.mkdir(parents=True)

    report = {
        "created_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "app_dir": str(APP_DIR),
        "threshold": {
            "model_load_seconds_max": 180,
            "model_generation_seconds_max": 60,
            "web_fetched_pages_min": 1,
            "pptx_slides_min": 4,
            "pptx_images_min": 1,
            "pptx_hyperlinks_min": 2,
        },
        "probes": [],
    }
    brain = KiraBrain()
    original_workspace = brain.agentic_workspace
    original_index = brain.artifacts_index_path
    original_attachments = brain.attachments_workspace
    try:
        brain.agentic_workspace = str(PROBE_WORKSPACE)
        brain.artifacts_index_path = str(PROBE_WORKSPACE / "artifact_index.json")
        brain.attachments_workspace = str(PROBE_WORKSPACE / "attachments")
        Path(brain.attachments_workspace).mkdir(parents=True, exist_ok=True)

        report["probes"].append(timed("model_load", lambda: (
            brain._evict_and_load("orchestrator")
            or {
                "ok": brain.current_brain == "orchestrator",
                "model_path": brain.orch_path,
                "worker_pid": brain.model_worker_process.pid if brain.model_worker_process else None,
            }
        )))

        def model_generation():
            prompt = brain._tokenizer_prompt([{
                "role": "user",
                "content": (
                    "You are Orchestrator V1. This is a release health check. "
                    "Reply with one concise sentence containing the exact marker KIRA_MODEL_READY."
                ),
            }])
            raw = brain._generate_sync(prompt, temperature=0.1, max_tokens=64, timeout_seconds=60)
            public = brain._strip_private_reasoning(raw)
            return {
                "ok": "KIRA_MODEL_READY" in raw and bool(public.strip()),
                "raw_tail": raw[-1200:],
                "public": public[-800:],
            }

        report["probes"].append(timed("model_generation", model_generation))

        def web_research():
            result = brain._web_research_tool(
                "Python official documentation language reference",
                max_pages="2",
            )
            fetched = 0
            import re
            match = re.search(r"Fetched pages:\s*(\d+)", result)
            if match:
                fetched = int(match.group(1))
            network_blocked = fetched == 0 and any(marker in result.lower() for marker in (
                "network is unreachable",
                "operation not permitted",
                "name or service not known",
                "temporary failure in name resolution",
                "no search provider returned usable result links",
            ))
            return {
                "ok": fetched >= 1 and "python" in result.lower(),
                "status": "blocked" if network_blocked else ("passed" if fetched >= 1 else "failed"),
                "fetched_pages": fetched,
                "result_excerpt": result[:5000],
            }

        report["probes"].append(timed("web_research", web_research))
        report["probes"].append(timed("artifacts", lambda: artifact_probe(brain)))

        def attachment():
            payload = base64.b64encode(b"KIRA ultra release attachment evidence").decode("ascii")
            imported = brain.import_attachment("probe-notes.md", payload, "text/markdown", "ultra-probe")
            normalized = brain._normalize_task_attachments([imported], "ultra-probe")
            context = brain._build_attachment_context(normalized)
            return {
                "ok": imported.get("ok", False) and "ultra release attachment evidence" in context,
                "imported": imported,
                "context_excerpt": context[:1200],
            }

        report["probes"].append(timed("attachment_ingestion", attachment))
    finally:
        brain.agentic_workspace = original_workspace
        brain.artifacts_index_path = original_index
        brain.attachments_workspace = original_attachments
        brain.is_running = False
        brain._stop_model_worker()

    for item in report["probes"]:
        detail = item.get("detail")
        if isinstance(detail, dict) and detail.get("status") == "blocked":
            item["status"] = "blocked"
    statuses = [item.get("status", "passed" if item.get("ok") else "failed") for item in report["probes"]]
    report["summary"] = {
        "passed": statuses.count("passed"),
        "failed": statuses.count("failed"),
        "blocked": statuses.count("blocked"),
    }
    report["ok"] = all(status == "passed" for status in statuses)
    report["inconclusive"] = not report["ok"] and "failed" not in statuses and "blocked" in statuses
    REPORT_PATH.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({
        "ok": report["ok"],
        "report": str(REPORT_PATH),
        "probes": [
            {
                "name": item["name"],
                "ok": item["ok"],
                "status": item.get("status", "passed" if item["ok"] else "failed"),
                "elapsed_seconds": item["elapsed_seconds"],
            }
            for item in report["probes"]
        ],
    }, indent=2))
    return 0 if report["ok"] else (2 if report["inconclusive"] else 1)


if __name__ == "__main__":
    raise SystemExit(main())
