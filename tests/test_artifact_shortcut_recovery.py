import json
import unittest
from unittest.mock import Mock, patch

import test_tree_web_artifacts as fixtures


class ArtifactRoutingTests(unittest.TestCase):
    setUp = fixtures.TreeWebArtifactTests.setUp
    tearDown = fixtures.TreeWebArtifactTests.tearDown

    def prepare_model(self, responses):
        self.actions.current_brain = "orchestrator"
        self.actions.active_model = object()
        self.actions._tokenizer_prompt = lambda messages: json.dumps(messages)
        self.actions._generate_with_watchdog = Mock(side_effect=responses)
        self.actions._split_orchestrator_response = lambda value: ("", value)
        self.actions._log_private_thoughts = Mock()

    def test_screenshot_prompt_recovers_scoped_research_without_new_task(self):
        prompt = "hey t can ypu build me a presentation ppt on the best anime ever made"
        blocks = self.actions._model_research_shortcut_blocks(
            '/quick web-research-brief "unrelated instructions"', prompt, "test-chat"
        )
        self.assertIn("[WEB_RESEARCH]", blocks)
        self.assertIn("QUERY: the best anime ever made", blocks)
        self.assertNotIn("unrelated", blocks)
        self.assertTrue(self.actions._prompt_needs_web_evidence_for_artifact(prompt))
        self.assertEqual(self.actions._prompt_keywords_for_artifact(prompt), ["anime"])

    def test_shortcut_recovery_does_not_execute_arbitrary_commands_or_offline_research(self):
        for output, prompt in [
            ("/quick delete-files", "create a ppt on anime"),
            ("Try /quick web-research-brief", "create a ppt on anime"),
            ("/quick web-research-brief", "hello"),
            ("/quick web-research-brief", "create an offline ppt on anime"),
        ]:
            self.assertEqual(self.actions._model_research_shortcut_blocks(output, prompt), "")
        self.assertFalse(self.actions._prompt_needs_web_evidence_for_artifact("offline ppt on the best anime"))
        self.assertFalse(self.actions._prompt_needs_web_evidence_for_artifact("create a report using my best writing style"))

    def test_quoted_source_path_does_not_request_deletion(self):
        blocks = self.actions._build_deterministic_agentic_blocks("create a ppt from `/tmp/source.pdf`")
        self.assertNotIn("DELETE_PATH", blocks)
        self.assertIn("DELETE_PATH", self.actions._build_deterministic_agentic_blocks("delete `/tmp/source.pdf`"))

    def test_research_then_artifact_handoff_for_each_format(self):
        for kind, tag in [("pptx", "NATIVE_PPTX"), ("docx", "NATIVE_DOCX"), ("pdf", "NATIVE_PDF")]:
            with self.subTest(kind=kind):
                research = "[WEB_RESEARCH]\nQUERY: robotics\n[/WEB_RESEARCH]"
                artifact = f"[{tag}]\nTITLE: Robotics\n[/{tag}]"
                self.prepare_model([research, artifact])
                evidence = "WEB_RESEARCH:\nFetched pages: 1\nReadable text: verified robotics reference"
                result = f"Generated: `/tmp/robotics.{kind}`"
                runner = Mock(side_effect=[evidence, result])
                self.actions._run_agentic_capabilities_with_watchdog = runner
                actual = self.actions._generate_prompt_specific_artifact_with_orchestrator(
                    kind, f"create a {kind} about robotics", "", "", "test-chat"
                )
                self.assertEqual(actual, result)
                self.assertEqual(runner.call_count, 2)
                self.assertEqual(runner.call_args_list[1].args[0], artifact)
                second_prompt = self.actions._generate_with_watchdog.call_args_list[1].args[0]
                self.assertIn("verified robotics reference", second_prompt)
                self.assertEqual(self.actions.web_evidence_by_chat["test-chat"], evidence)

    def test_repeated_shortcut_is_bounded_and_cached_research_is_not_repeated(self):
        self.prepare_model(["/quick web-research-brief"] * 2)
        runner = Mock()
        self.actions._run_agentic_capabilities_with_watchdog = runner
        result = self.actions._generate_prompt_specific_artifact_with_orchestrator(
            "pptx", "create a ppt on anime", "", "WEB_RESEARCH:\nFetched pages: 1", "test-chat"
        )
        self.assertEqual(result, "")
        self.assertEqual(self.actions._generate_with_watchdog.call_count, 2)
        runner.assert_not_called()

    def test_failed_research_does_not_produce_artifact_or_claim_success(self):
        self.prepare_model(["/quick web-research-brief"])
        self.actions._run_agentic_capabilities_with_watchdog = Mock(return_value="WEB_RESEARCH:\nFetched pages: 0")
        result = self.actions._generate_prompt_specific_artifact_with_orchestrator(
            "docx", "create a document on robotics", "", "", "test-chat"
        )
        self.assertEqual(result, "")
        self.assertEqual(self.actions._generate_with_watchdog.call_count, 1)

    def test_artifact_repair_does_not_execute_unrelated_mutations(self):
        artifact = "[NATIVE_DOCX]\nTITLE: Robotics\n[/NATIVE_DOCX]"
        self.prepare_model(["[DELETE_PATH]\nPATH: /tmp/important\n[/DELETE_PATH]\n" + artifact])
        runner = Mock(return_value="DOCX generated")
        self.actions._run_agentic_capabilities_with_watchdog = runner
        self.actions._generate_prompt_specific_artifact_with_orchestrator("docx", "create a document on robotics", "", "", "test-chat")
        self.assertEqual(runner.call_args.args[0], artifact)

    def test_recovered_shortcut_creates_verified_pptx_with_real_slides_agent(self):
        from pathlib import Path
        from pptx import Presentation

        prompt = "hey t can ypu build me a presentation ppt on the best anime ever made and open it"
        sections = [
            ("Animation", "Anime uses drawing, timing, composition and sound together. The visual language ranges from restrained character acting to highly stylized motion, so animation quality should be evaluated in relation to each work's aims."),
            ("Story", "Long-form anime can develop relationships across many episodes, while films often rely on a tighter narrative structure. A useful comparison considers pacing, character motivations and whether the ending follows from earlier choices."),
            ("Sound", "Music, voice acting and sound design shape the experience of an anime. Strong sound can establish atmosphere or emphasize a character's emotional state, while silence can give important scenes room to unfold."),
            ("Audience", "There is no single best anime for every viewer. Preferences differ across genres and age groups, so a recommendation should explain its selection criteria instead of treating popularity alone as proof of artistic quality."),
            ("Choosing", "Compare an anime's storytelling, visual direction, sound and suitability for the intended viewer. State whether a ranking reflects personal taste, critical reception or audience ratings, since these measures answer different questions."),
        ]
        spec = "[NATIVE_PPTX]\nTITLE: Anime: Comparing the Art Form\n" + "\n".join(
            f"SLIDE: {title}\nBODY: {body}" for title, body in sections
        ) + "\n[/NATIVE_PPTX]"
        self.prepare_model([spec])
        self.actions._open_target_tool = Mock(side_effect=lambda path: f"Opened: {path}")
        calls = []

        def run(blocks, request, chat_id, label):
            calls.append(label)
            if "[WEB_RESEARCH]" in blocks:
                return "WEB_RESEARCH:\nFetched pages: 1\nReadable text: Fixture evidence about anime comparison criteria."
            return self.actions._generate_native_pptx_blocks(blocks, request, chat_id)

        self.actions._run_agentic_capabilities_with_watchdog = run
        result = self.actions._ensure_requested_artifact_or_real_open(
            prompt, '/quick web-research-brief "acclaimed anime"', "", "test-chat"
        )
        path = self.actions._extract_artifact_path_from_text(result)
        self.assertTrue(path and Path(path).is_file(), result)
        self.assertTrue(self.actions._pptx_quality_report(path, prompt)["ok"], result)
        self.assertGreaterEqual(len(Presentation(path).slides), 5)
        self.actions._open_target_tool.assert_called_once_with(path)
        self.assertEqual(calls, ["artifact_web_evidence", "artifact_repair"])
        self.assertNotIn("/quick", result)

    def test_ui_retains_wordmark_without_logo(self):
        from test_tree_web_artifacts import APP_DIR
        ui = (APP_DIR / "ui.html").read_text()
        self.assertNotIn('class="kira-logo-image"', ui)
        self.assertIn('class="kira-brand-product">SUPERAPP', ui)

    def test_artifact_reloads_worker_after_timeout(self):
        self.prepare_model(['[NATIVE_DOCX]\nTITLE: Science\n[/NATIVE_DOCX]'])
        self.actions.active_model = None
        self.actions.current_brain = None
        self.actions._evict_and_load = Mock()
        self.actions._run_agentic_capabilities_with_watchdog = Mock(return_value='executed')
        result = self.actions._generate_prompt_specific_artifact_with_orchestrator(
            'docx', 'Create a document about science', '', '', 'test-chat'
        )
        self.actions._evict_and_load.assert_called_once_with('orchestrator')
        self.assertEqual(result, 'executed')
        prompt = self.actions._generate_with_watchdog.call_args.args[0]
        self.assertFalse(prompt.endswith('<thought_process>\n'))
        self.assertIn('Always include the closing tag', prompt)

    def test_science_request_uses_focused_author_and_creates_real_pptx(self):
        import os
        import queue
        from pathlib import Path
        from pptx import Presentation
        import interface

        request = 'Create a PPT on the different forms of science in our daily life.'
        topics = [
            ('Physics', 'Physics explains forces and energy in daily life. Walking depends on friction between shoes and the ground. A bicycle converts muscular effort into motion, while its brakes convert kinetic energy into heat. These examples connect ordinary activities with measurable physical principles.'),
            ('Chemistry', 'Chemistry describes substances and how they change. Cooking transforms ingredients through chemical reactions, while soap helps water remove oily dirt. Baking powder releases gas that makes cakes rise. Understanding these changes helps explain familiar results in kitchens and household cleaning.'),
            ('Biology', 'Biology studies living things and their interactions. Digestion breaks food into nutrients that cells can use. Plants capture light energy through photosynthesis, supporting food chains. Gardening, nutrition and hygiene all draw on biological knowledge to explain growth, health and living environments.'),
            ('Earth Science', 'Earth science connects daily life with weather, rocks and water. Evaporation, condensation and precipitation move water through the environment. Weather forecasts help people plan outdoor activities. Soil properties influence gardening, and knowledge of natural hazards supports preparation and safer choices.'),
            ('Astronomy', 'Astronomy studies objects and processes beyond Earth. Earth rotates to produce day and night, while its tilted axis causes seasons during its orbit around the Sun. Observing the sky connects calendars, seasonal patterns and daily routines with larger astronomical cycles.'),
        ]
        spec = '[NATIVE_PPTX]\nTITLE: Science in Our Daily Life\n' + '\n'.join(
            f'SLIDE: {title}\nBODY: {body}' for title, body in topics
        ) + '\n[/NATIVE_PPTX]'
        with patch.dict(os.environ, {'HOME': str(self.root)}), patch.object(interface, 'BASE_DIR', str(self.root)), patch('threading.Thread.start'):
            brain = interface.KiraBrain()
        brain._evict_and_load = Mock()
        brain.current_brain = 'orchestrator'
        brain.active_model = object()
        brain._generate_with_watchdog = Mock(return_value=spec)
        brain._run_orchestrator_agent_loop = Mock(side_effect=AssertionError('General loop should not run'))
        brain._build_system_instruction = Mock(side_effect=AssertionError('General tool catalog should not be built'))
        brain._smart_memory_context = Mock(return_value='')
        brain._maybe_update_personalization_rag = Mock()
        brain._route_and_generate({'prompt': request, 'mode': 'chat'})
        updates = []
        while not brain.response_queue.empty():
            updates.append(brain.response_queue.get_nowait())
        errors = [u for u in updates if u['type'] == 'error']
        self.assertEqual(errors, [])
        messages = [u['content'] for u in updates if u['type'] == 'message']
        self.assertTrue(messages, updates)
        path = brain._extract_artifact_path_from_text(messages[-1])
        self.assertTrue(path and Path(path).is_file(), messages)
        self.assertGreaterEqual(len(Presentation(path).slides), 5)
        self.assertTrue(brain._pptx_quality_report(path, request)['ok'])
        brain._run_orchestrator_agent_loop.assert_not_called()
        brain._build_system_instruction.assert_not_called()
        brain._smart_memory_context.assert_called_once()
        self.assertEqual(brain._generate_with_watchdog.call_count, 1)


if __name__ == "__main__":
    unittest.main()
