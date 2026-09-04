import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from interface import KiraBrain


class ClipboardAttachmentTests(unittest.TestCase):
    def test_native_urls_and_legacy_filenames(self):
        import AppKit
        url = SimpleNamespace(isFileURL=lambda: True, path=lambda: "/tmp/copied document.pdf")
        board = SimpleNamespace(
            readObjectsForClasses_options_=lambda *args: [url, url],
            propertyListForType_=lambda kind: ["/tmp/legacy.docx"],
        )
        brain = KiraBrain.__new__(KiraBrain)
        with patch("interface.sys.platform", "darwin"), patch.object(
            AppKit, "NSPasteboard", SimpleNamespace(generalPasteboard=lambda: board)
        ):
            self.assertEqual(brain._native_clipboard_file_paths(), ["/tmp/copied document.pdf"])
            board.readObjectsForClasses_options_ = lambda *args: []
            self.assertEqual(brain._native_clipboard_file_paths(), ["/tmp/legacy.docx"])

    def test_copied_documents_import_without_base64_and_keep_original(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as root:
            brain = KiraBrain.__new__(KiraBrain)
            brain.attachments_workspace = os.path.join(root, "attachments")
            source = Path(root) / "copied report.docx"
            document = Document()
            document.add_paragraph("Native clipboard document evidence")
            document.save(source)
            original = source.read_bytes()
            brain._native_clipboard_file_paths = lambda: [str(source), str(source)]
            result = brain.paste_attachments("chat-1")
            self.assertTrue(result["ok"])
            self.assertEqual(len(result["files"]), 1)
            item = result["files"][0]
            self.assertTrue(item["ok"])
            self.assertNotEqual(item["path"], str(source))
            self.assertEqual(source.read_bytes(), original)
            self.assertIn("Native clipboard document evidence", brain._attachment_text(item))

    def test_bad_file_does_not_abort_remaining_files(self):
        with tempfile.TemporaryDirectory() as root:
            brain = KiraBrain.__new__(KiraBrain)
            brain.attachments_workspace = os.path.join(root, "attachments")
            good = Path(root) / "notes.txt"
            good.write_text("notes")
            bad = Path(root) / "payload.exe"
            bad.write_text("executable")
            result = brain._import_attachment_paths([root, str(bad), str(good)], "chat")
            self.assertEqual([item["ok"] for item in result["files"]], [False, False, True])

    def test_empty_clipboard_does_not_import_plain_text_as_a_path(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain._native_clipboard_file_paths = lambda: []
        self.assertEqual(brain.paste_attachments("chat"), {"ok": True, "files": []})

    def test_attachment_context_is_bounded_and_shares_space(self):
        brain = KiraBrain.__new__(KiraBrain)
        seen = []
        def extract(item, max_chars=45000):
            seen.append(max_chars)
            return item["name"] + " evidence " * 10000
        brain._attachment_text = extract
        items = [{"name": f"doc{i}.txt", "path": f"/tmp/doc{i}.txt", "mime": "text/plain"} for i in range(8)]
        context = brain._build_attachment_context(items)
        self.assertLessEqual(len(context), 6000)
        for item in items:
            self.assertIn(item["name"], context)
        self.assertLess(max(seen), 1000)

    def test_pptx_attachment_extraction_uses_supported_iteration(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "slides.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[0])
            slide.shapes.title.text = "Slide evidence"
            presentation.save(path)
            text = KiraBrain.__new__(KiraBrain)._attachment_text({"path": str(path)}, max_chars=500)
            self.assertIn("Slide evidence", text)
            self.assertNotIn("Extraction warning", text)

    def test_spreadsheet_extraction_is_bounded(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "sheet.xlsx"
            book = Workbook()
            book.active.append(["Spreadsheet evidence"])
            book.active.cell(row=10000, column=100, value="OUTSIDE_EXCERPT")
            book.save(path)
            book.close()
            text = KiraBrain.__new__(KiraBrain)._attachment_text({"path": str(path)}, max_chars=500)
            self.assertIn("Spreadsheet evidence", text)
            self.assertNotIn("OUTSIDE_EXCERPT", text)
            self.assertLessEqual(len(text), 500)
