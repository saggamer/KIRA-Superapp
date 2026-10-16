import queue
import sys
import tempfile
import types
import unittest
from pathlib import Path

if 'webview' not in sys.modules:
    sys.modules['webview'] = types.SimpleNamespace()

from interface import KiraBrain
from actions import KiraActionsMixin
from orchestrator_v15.adapter_router import channel_for_request
from orchestrator_v15.presentation_controller import PresentationController, is_presentation_request
from scripts.analyze_image_with_gemma4 import visual_profile


def deck(slides=5):
    body = ['[NATIVE_PPTX]', 'TITLE: Battery Recycling Brief', 'THEME: light blue']
    for index in range(slides):
        body += [f'SLIDE: Finding {index + 1}',
                 f'BODY: The supplied brief describes finding {index + 1} and its implication for the audience.']
    body.append('[/NATIVE_PPTX]')
    return '\n'.join(body)


class PresentationControllerTests(unittest.TestCase):
    def test_supplied_local_images_do_not_require_web_evidence(self):
        actions = KiraActionsMixin.__new__(KiraActionsMixin)
        self.assertFalse(actions._prompt_needs_web_evidence_for_artifact(
            'Create a PowerPoint using the provided image /tmp/kira-desktop.png.'))
        self.assertTrue(actions._prompt_needs_web_evidence_for_artifact(
            'Create a PowerPoint and find relevant images on the web.'))

    def test_visual_profile_reports_dimensions_and_readable_theme(self):
        from PIL import Image, ImageDraw
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / 'slide-source.png'
            image = Image.new('RGB', (320, 180), '#f9fafb')
            ImageDraw.Draw(image).rectangle((20, 20, 110, 150), fill='#238aaf')
            image.save(path)
            profile = visual_profile(path)
        self.assertIn('Dimensions: 320x180', profile)
        self.assertIn('Suggested THEME: light blue', profile)

    def test_image_analysis_accepts_model_path_and_task_aliases(self):
        path, question = KiraActionsMixin._image_analysis_request({
            'target': '`/workspace/chart.png`',
            'task': 'Read the chart and identify an audience-fit visual direction.',
        })
        self.assertEqual(path, '/workspace/chart.png')
        self.assertIn('audience-fit', question)

    def test_deck_about_browser_routes_to_control(self):
        self.assertTrue(is_presentation_request('Create a PowerPoint about browser safety'))
        self.assertEqual(channel_for_request('Create a PowerPoint about browser safety'), 'control')

    def test_requires_source_for_current_claims_and_respects_supplied_context(self):
        controller = PresentationController('Create a presentation on the latest battery policy')
        self.assertIn('WEB_RESEARCH', controller.preflight(deck()))
        controller.record('[WEB_RESEARCH]\nQUERY: latest battery policy\n[/WEB_RESEARCH]',
                          'Research result: official policy page was retrieved.')
        self.assertIsNone(controller.preflight(deck()))
        supplied = PresentationController('Create a presentation on the latest battery policy',
                                          context='Source: https://example.org/policy')
        self.assertIsNone(supplied.preflight(deck()))

    def test_slide_count_and_verification_are_bounded(self):
        controller = PresentationController('Create a 6 slide PowerPoint')
        self.assertIn('5 content slides', controller.preflight(deck(4)))
        self.assertIsNone(controller.preflight(deck(5)))
        self.assertIn('No verified', controller.preflight('Done.'))
        controller.record(deck(5), 'PPTX generated: `/tmp/deck.pptx`\nQuality verified: 6 slides, 0 images.')
        self.assertIsNone(controller.preflight('Done.'))

    def test_content_slide_count_can_be_spelled_out(self):
        controller = PresentationController('Make five substantive content slides plus a cover.')
        self.assertEqual(controller.min_slides, 5)
        self.assertIn('5 content slides', controller.preflight(deck(4)))

    def test_research_complete_gives_specific_malformed_deck_repair(self):
        controller = PresentationController('Create a PowerPoint with sources about a fictional topic.')
        controller.record('[WEB_RESEARCH]\nQUERY: fictional topic source\n[/WEB_RESEARCH]',
                          'WEB_RESEARCH:\nFetched pages: 1\nURL: https://example.org/brief')
        error = controller.preflight('I have a plan.\n[/NATIVE_PPTX]\nTITLE: Fictional Topic')
        self.assertIn('Start the response with [NATIVE_PPTX]', error)
        self.assertIn('Research is complete', controller.result_hint())

    def test_visual_deck_requires_image_evidence(self):
        controller = PresentationController('Create a PowerPoint with images about the library.')
        self.assertIn('WEB_IMAGE_SEARCH', controller.preflight(deck()))
        self.assertIn('WEB_IMAGE_SEARCH', controller.preflight(
            '[WEB_SEARCH]\nQUERY: library pictures\n[/WEB_SEARCH]'))
        offline = PresentationController('Create a text-only PowerPoint with no images.')
        self.assertFalse(offline.needs_images)
        brain = KiraBrain.__new__(KiraBrain)
        self.assertFalse(brain._pptx_visuals_requested('Create a text-only PowerPoint with no images.'))

    def test_number_only_slide_heading_is_rejected(self):
        controller = PresentationController('Create a PowerPoint about a fictional library.')
        numbered = deck().replace('SLIDE: Finding 1', 'SLIDE: 1/6')
        self.assertIn('descriptive title', controller.preflight(numbered))

    def test_model_cover_is_rejected_before_rendering(self):
        controller = PresentationController('Create a PowerPoint about a fictional library.')
        proposed = deck().replace('SLIDE: Finding 1', 'SLIDE: 1/5 - Cover')
        self.assertIn('creates the cover automatically', controller.preflight(proposed))

    def test_renderer_uses_model_cover_image_without_duplicate_cover(self):
        from PIL import Image
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as temporary:
            image_path = Path(temporary) / 'cover.png'
            Image.new('RGB', (640, 360), '#dc7c22').save(image_path)
            actions = KiraActionsMixin.__new__(KiraActionsMixin)
            actions.agentic_workspace = temporary
            actions._record_artifact = lambda *_args: None
            actions._log_agentic_event = lambda *_args: None
            actions._materialize_pptx_image = lambda path, _title: (path, '')
            block = ('TITLE: Fictional Library\nTHEME: light orange\n'
                     'SLIDE: 1/3 - Cover\nBODY: Welcome\nIMAGE: ' + str(image_path) + '\n'
                     'SLIDE: 2/3 - Reading Room\nBODY: A quiet room.\n'
                     'SLIDE: 3/3 - Events\nBODY: Weekly talks.')
            result = actions._generate_native_pptx(block, record_artifact=False)
            self.assertIn('PPTX generated:', result)
            path = Path(result.split('`')[1])
            slides = Presentation(path).slides
            self.assertEqual(len(slides), 3)
            self.assertEqual(sum(shape.shape_type == 13 for shape in slides[0].shapes), 1)
            self.assertIn('Reading Room', '\n'.join(shape.text for shape in slides[1].shapes if shape.has_text_frame))
            self.assertNotIn('2/3', '\n'.join(shape.text for shape in slides[1].shapes if shape.has_text_frame))

    def test_explicit_web_first_order_precedes_image_search(self):
        controller = PresentationController(
            'Create a PowerPoint with images; use web search first and cite sources.')
        image_search = '[WEB_IMAGE_SEARCH]\nQUERY: official visual\n[/WEB_IMAGE_SEARCH]'
        self.assertIn('WEB_RESEARCH', controller.preflight(image_search))
        controller.record('[WEB_RESEARCH]\nQUERY: official story\n[/WEB_RESEARCH]',
                          'WEB_RESEARCH:\nFetched pages: 1\nURL: https://example.org/story')
        self.assertIsNone(controller.preflight(image_search))

    def test_image_analysis_is_required_before_image_led_deck(self):
        request = 'Create a PowerPoint from /workspace/chart.png for the board.'
        controller = PresentationController(request)
        analyze = ('[IMAGE_ANALYZE]\nPATH: /workspace/chart.png\n'
                   'QUESTION: Read values and suggest a clear story.\n[/IMAGE_ANALYZE]')
        self.assertIsNone(controller.preflight(analyze))
        self.assertIn('Analyze at least one', controller.preflight(deck()))
        controller.record(analyze, 'IMAGE_ANALYSIS:\nObserved chart values and design suggestion.')
        self.assertIsNone(controller.preflight(deck()))

    def test_downloaded_image_path_with_spaces_is_checked_as_one_path(self):
        controller = PresentationController('Create a PowerPoint with one image.')
        controller.evidence_text += '\nPath: `/tmp/Kira slide image.jpg`'
        action = '[IMAGE_ANALYZE]\nPATH: /tmp/Kira slide image.jpg\n[/IMAGE_ANALYZE]'
        self.assertIsNone(controller.preflight(action))
        controller.record(action, 'IMAGE_ANALYSIS:\nThe image is readable.')
        self.assertIn('/tmp/Kira slide image.jpg', controller.analyzed_image_paths)

    def test_remote_image_candidates_do_not_force_visual_analysis(self):
        controller = PresentationController(
            'Create a source-linked PowerPoint about a fictional topic.',
            context='WEB_RESEARCH:\nURL: https://example.org/brief\n'
                    'Image candidates:\n- https://example.org/candidate.png')
        controller.research_observed = True
        self.assertFalse(controller.context_has_images)
        with_link = deck().replace('[/NATIVE_PPTX]',
                                   'LINK: Source brief | https://example.org/brief\n[/NATIVE_PPTX]')
        self.assertIsNone(controller.preflight(with_link))

    def test_visual_evidence_requires_download_and_successful_analysis(self):
        controller = PresentationController('Create a PowerPoint with images about the library.')
        search = '[WEB_IMAGE_SEARCH]\nQUERY: library reading room photos\n[/WEB_IMAGE_SEARCH]'
        controller.record(search, 'WEB_IMAGE_SEARCH:\nNo images downloaded. Candidate URLs:\n- https://example.org/room.jpg')
        self.assertFalse(controller.image_research_observed)
        download = '[WEB_IMAGE_DOWNLOAD]\nURL: https://example.org/room.jpg\n[/WEB_IMAGE_DOWNLOAD]'
        self.assertIsNone(controller.preflight(download))
        controller.record(download, 'WEB_IMAGE_DOWNLOAD:\nSource: `https://example.org/room.jpg`\nPath: `/tmp/room.jpg`')
        self.assertTrue(controller.image_research_observed)
        mistaken_download = '[WEB_IMAGE_DOWNLOAD]\nIMAGE_PATH: /tmp/room.jpg\n[/WEB_IMAGE_DOWNLOAD]'
        self.assertIn('already a downloaded local file', controller.preflight(mistaken_download))
        analyze = '[IMAGE_ANALYZE]\nTARGET: /tmp/room.jpg\nTASK: Identify palette and layout.\n[/IMAGE_ANALYZE]'
        self.assertIsNone(controller.preflight(analyze))
        controller.record(analyze, 'IMAGE_ANALYZE error: image decoder unavailable')
        self.assertFalse(controller.image_analysis_observed)
        controller.record(analyze, 'IMAGE_ANALYSIS:\nPath: `/tmp/room.jpg`\nPalette: warm neutrals. Composition: bright subject.')
        self.assertTrue(controller.image_analysis_observed)
        self.assertIn('2 distinct local image', controller.preflight(deck()))
        second = '[WEB_IMAGE_DOWNLOAD]\nURL: https://example.org/room2.jpg\n[/WEB_IMAGE_DOWNLOAD]'
        controller.record(search, 'WEB_IMAGE_SEARCH:\nCandidate URLs:\n- https://example.org/room2.jpg')
        self.assertIsNone(controller.preflight(second))
        controller.record(second, 'WEB_IMAGE_DOWNLOAD:\nSource: `https://example.org/room2.jpg`\nPath: `/tmp/room2.jpg`')
        analyze_second = '[IMAGE_ANALYZE]\nPATH: /tmp/room2.jpg\n[/IMAGE_ANALYZE]'
        controller.record(analyze_second, 'IMAGE_ANALYSIS:\nPath: `/tmp/room2.jpg`\nReadable slide image.')
        with_images = deck().replace('[/NATIVE_PPTX]', 'IMAGE: /tmp/room.jpg\nIMAGE: /tmp/room2.jpg\n[/NATIVE_PPTX]')
        self.assertIsNone(controller.preflight(with_images))
        self.assertIn('2 distinct grounded IMAGE', controller.preflight(
            with_images.replace('IMAGE: /tmp/room2.jpg', 'IMAGE: /tmp/room.jpg')))
        self.assertIn('THEME', controller.preflight(with_images.replace('THEME: light blue\n', '')))

    def test_reads_supplied_file_before_source_backed_deck(self):
        controller = PresentationController('Create a presentation from /workspace/brief.md with current figures.')
        read = '[READ_FILE]\nPATH: /workspace/brief.md\n[/READ_FILE]'
        self.assertIsNone(controller.preflight(read))
        self.assertIn('evidence', controller.preflight(deck()))
        controller.record(read, 'READ_FILE:\nThe brief lists current figures and their source.')
        self.assertIsNone(controller.preflight(deck()))

    def test_rejects_invented_slide_source(self):
        controller = PresentationController('Create a PowerPoint with sources about the library.',
                                            context='Source: https://example.org/brief')
        sourced = deck().replace('[/NATIVE_PPTX]',
                                'LINK: Policy | https://unseen.example/policy\n[/NATIVE_PPTX]')
        self.assertIn('not in', controller.preflight(sourced))

    def test_accepts_direct_source_url_without_label(self):
        controller = PresentationController('Create a PowerPoint with sources about the library.',
                                            context='Source: https://example.org/brief')
        sourced = deck().replace('[/NATIVE_PPTX]',
                                'LINK: https://example.org/brief\n[/NATIVE_PPTX]')
        self.assertIsNone(controller.preflight(sourced))

    def test_diagonal_profile_executes_only_selected_deck(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain.response_queue = queue.Queue()
        brain._evict_and_load = lambda _model: None
        brain._truncate = lambda text, length=1000: str(text)[:length]
        brain._load_chat_messages = lambda _chat_id: []
        brain._extract_agentic_blocks = lambda _text: []
        calls = []
        payloads = []

        def model_reply(payload, timeout=60):
            payloads.append(payload)
            if len(payloads) == 1:
                return {'proposals': [deck()], 'token_counts': [300], 'memo_tokens': 0}
            return {'proposals': ['The deck is ready.'], 'token_counts': [8], 'memo_tokens': 0}

        brain._model_worker_request = model_reply
        brain._run_agentic_capabilities_with_watchdog = lambda answer, *_: (
            calls.append(answer) or 'PPTX generated: `/tmp/deck.pptx`\nQuality verified: 6 slides, 0 images.')
        answer = brain._run_v15_profile('Create a PowerPoint about battery recycling.', 'chat-ppt', 'diagonal')
        self.assertIn('/tmp/deck.pptx', answer)
        self.assertEqual(calls, [deck()])
        self.assertTrue(payloads[0]['presentation_task'])

    def test_diagonal_researches_before_creating_current_deck(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain.response_queue = queue.Queue()
        brain._evict_and_load = lambda _model: None
        brain._truncate = lambda text, length=1000: str(text)[:length]
        brain._load_chat_messages = lambda _chat_id: []
        brain._extract_agentic_blocks = lambda _text: []
        calls = []
        research = '[WEB_RESEARCH]\nQUERY: latest battery recycling changes\n[/WEB_RESEARCH]'
        outputs = [
            {'proposals': [deck(), research, 'Finished.'], 'token_counts': [250, 20, 5]},
            {'proposals': [deck()], 'token_counts': [300]},
            {'proposals': ['The deck is ready.'], 'token_counts': [8]},
        ]
        seen = []

        def model_reply(payload, timeout=60):
            seen.append(payload)
            return outputs[len(seen) - 1]

        def execute(answer, *_):
            calls.append(answer)
            return ('Research result: current data retrieved.' if answer == research else
                    'PPTX generated: `/tmp/current-deck.pptx`\nQuality verified: 6 slides, 0 images.')

        brain._model_worker_request = model_reply
        brain._run_agentic_capabilities_with_watchdog = execute
        result = brain._run_v15_profile('Create a PowerPoint on the latest battery recycling changes.',
                                        'chat-ppt-current', 'diagonal')
        self.assertIn('/tmp/current-deck.pptx', result)
        self.assertEqual(calls, [research, deck()])
        self.assertEqual(seen[0]['branches'], 3)
        self.assertEqual(seen[1]['branches'], 3)

    def test_research_evidence_reaches_slide_specialist(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain.response_queue = queue.Queue()
        brain._evict_and_load = lambda _model: None
        brain._truncate = lambda text, length=1000: str(text)[:length]
        brain._load_chat_messages = lambda _chat_id: []
        brain._extract_agentic_blocks = lambda _text: []
        brain.web_evidence_by_chat = {}
        research = '[WEB_RESEARCH]\nQUERY: latest battery recycling changes\n[/WEB_RESEARCH]'
        replies = [research, deck(), 'The deck is ready.']
        brain._model_worker_request = lambda *_args, **_kwargs: {
            'proposals': [replies.pop(0)], 'token_counts': [30], 'memo_tokens': 0}

        def execute(answer, *_):
            if answer == research:
                return 'Fetched pages: 1\nSource: https://example.org/brief\nBattery recycling grew.'
            self.assertIn('https://example.org/brief', brain.web_evidence_by_chat['chat-ppt-evidence'])
            return 'PPTX generated: `/tmp/deck.pptx`\nQuality verified: 6 slides, 0 images.'

        brain._run_agentic_capabilities_with_watchdog = execute
        answer = brain._run_v15_profile('Create a PowerPoint on the latest battery recycling changes.',
                                        'chat-ppt-evidence', 'diagonal')
        self.assertIn('/tmp/deck.pptx', answer)


if __name__ == '__main__':
    unittest.main()
