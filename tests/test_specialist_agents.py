import os
import sys
import tempfile
import time
import unittest
from pathlib import Path


APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

from actions import KiraActionsMixin


class SpecialistHarness(KiraActionsMixin):
    def __init__(self, root):
        self.app_root = str(root)
        self.home_path = str(root)
        self.agentic_workspace = str(root / "workspace")
        self.agentic_logs_path = str(root / "logs")
        self.architect_workspace = str(root / "architect")
        self.computer_roots = {"test": str(root)}
        self.readable_roots = [str(root)]
        self.web_evidence_by_chat = {}
        self.last_agentic_intent = None
        self.last_agentic_result = None
        os.makedirs(self.agentic_workspace, exist_ok=True)
        os.makedirs(self.agentic_logs_path, exist_ok=True)

    def _log_agentic_event(self, *_args, **_kwargs):
        return None

    def _attach_artifact_provenance(self, *_args, **_kwargs):
        return None


class SpecialistAgentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.actions = SpecialistHarness(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def test_registry_is_exposed_through_tree(self):
        registry = self.actions._specialist_registry_tool()
        tree = self.actions._branch_registry_tool()
        self.assertIn("slides: Slides Agent", registry)
        self.assertIn("artifact_verifier: Artifact Verifier", registry)
        self.assertIn("tree-4", (self.root / "architect" / "kira_tree_branch_registry.json").read_text())
        self.assertIn("specialists", tree)

    def test_research_and_media_specialists_run_in_parallel(self):
        from PIL import Image

        image_path = self.root / "robot-a.png"
        Image.effect_noise((640, 420), 48).convert("RGB").save(image_path)

        def delayed_research(_query, max_pages="5"):
            time.sleep(0.2)
            return (
                "WEB_RESEARCH:\nFetched pages: 2\n"
                "Source: https://example.com/robotics\n"
                "Readable text:\nverified robotics evidence"
            )

        def delayed_media(_query, limit="5"):
            time.sleep(0.2)
            return (
                "WEB_IMAGE_SEARCH:\nDownloaded images: 1\n"
                f"- Path: `{image_path}`\n"
                "  Source: https://example.com/robotics-image.png\n"
                "  Page: https://example.com/robotics\n"
                "  Alt: Robotics laboratory"
            )

        self.actions._web_research_tool = delayed_research
        self.actions._web_image_search_tool = delayed_media
        blocks = (
            "[SPECIALIST_TASK]\nAGENT: research\nQUERY: robotics\n[/SPECIALIST_TASK]\n"
            "[SPECIALIST_TASK]\nAGENT: media\nQUERY: robotics labs\n[/SPECIALIST_TASK]"
        )
        started = time.perf_counter()
        result = self.actions._run_specialist_tools(blocks, "chat")
        elapsed = time.perf_counter() - started
        self.assertLess(elapsed, 0.34)
        self.assertEqual(result.count("Status: verified"), 2)
        self.assertIn("Fetched pages: 2", result)
        self.assertIn("Downloaded images: 1", result)

    def test_media_agent_rejects_tiny_and_unrelated_assets(self):
        from PIL import Image

        tiny = self.root / "naruto-tiny.png"
        unrelated = self.root / "generic-template.png"
        Image.new("RGB", (264, 57), (128, 128, 128)).save(tiny)
        Image.effect_noise((800, 450), 48).convert("RGB").save(unrelated)

        self.actions._web_image_search_tool = lambda _query, limit="5": (
            "WEB_IMAGE_SEARCH:\nDownloaded images: 2\n"
            f"- Path: `{tiny}`\n"
            "  Source: https://naruto-official.com/tiny.png\n"
            "  Page: https://naruto-official.com/en/\n"
            "  Alt: Naruto\n"
            f"- Path: `{unrelated}`\n"
            "  Source: https://stock.example.com/mockup.png\n"
            "  Page: https://stock.example.com/templates\n"
            "  Alt: Instant mockup template"
        )
        result = self.actions._specialist_hub().execute({
            "agent": "media",
            "query": "Naruto Shippuden characters",
        }, "chat")
        self.assertEqual(result["status"], "failed", result)
        self.assertEqual(result["quality"]["accepted"], 0)
        self.assertEqual(result["quality"]["rejected"], 2)
        self.assertTrue(
            any("low resolution" in reason for reason in result["quality"]["rejection_reasons"])
        )
        self.assertTrue(
            any("does not match" in reason for reason in result["quality"]["rejection_reasons"])
        )

    def test_media_agent_accepts_relevant_presentation_asset(self):
        from PIL import Image

        image_path = self.root / "naruto-team-seven.jpg"
        Image.effect_noise((1200, 675), 48).convert("RGB").save(image_path, quality=92)
        self.actions._web_image_search_tool = lambda _query, limit="5": (
            "WEB_IMAGE_SEARCH:\nDownloaded images: 1\n"
            f"- Path: `{image_path}`\n"
            "  Source: https://example.com/naruto-team-seven.jpg\n"
            "  Page: https://example.com/naruto-shippuden\n"
            "  Alt: Naruto Shippuden Team Seven"
        )
        result = self.actions._specialist_hub().execute({
            "agent": "media",
            "query": "Naruto Shippuden characters",
        }, "chat")
        self.assertEqual(result["status"], "verified", result)
        self.assertEqual(result["assets"], [str(image_path)])
        self.assertEqual(result["quality"]["accepted"], 1)

    def test_media_agent_rejects_large_but_visually_blank_asset(self):
        from PIL import Image

        image_path = self.root / "naruto-placeholder.png"
        Image.new("RGB", (1280, 720), (35, 57, 81)).save(image_path)
        self.actions._web_image_search_tool = lambda _query, limit="5": (
            "WEB_IMAGE_SEARCH:\nDownloaded images: 1\n"
            f"- Path: `{image_path}`\n"
            "  Source: https://example.com/naruto-placeholder.png\n"
            "  Page: https://example.com/naruto-shippuden\n"
            "  Alt: Naruto Shippuden characters"
        )
        result = self.actions._specialist_hub().execute({
            "agent": "media",
            "query": "Naruto Shippuden characters",
        }, "chat")
        self.assertEqual(result["status"], "failed", result)
        self.assertTrue(any(
            "visually flat" in reason or "suspiciously small" in reason
            for reason in result["quality"]["rejection_reasons"]
        ), result)

    def test_media_agent_deduplicates_identical_visual_content(self):
        from PIL import Image

        first = self.root / "robotics-control-room-a.png"
        second = self.root / "robotics-control-room-b.png"
        Image.effect_noise((960, 540), 64).convert("RGB").save(first)
        second.write_bytes(first.read_bytes())
        self.actions._web_image_search_tool = lambda _query, limit="5": (
            "WEB_IMAGE_SEARCH:\nDownloaded images: 2\n"
            f"- Path: `{first}`\n"
            "  Source: https://example.com/robotics-a.png\n"
            "  Page: https://example.com/robotics\n"
            "  Alt: Robotics control room\n"
            f"- Path: `{second}`\n"
            "  Source: https://example.org/robotics-b.png\n"
            "  Page: https://example.org/robotics\n"
            "  Alt: Robotics control room"
        )

        result = self.actions._specialist_hub().execute({
            "agent": "media",
            "query": "robotics control room",
        }, "chat")

        self.assertEqual(result["status"], "verified", result)
        self.assertEqual(result["quality"]["accepted"], 1)
        self.assertEqual(result["quality"]["rejected"], 1)
        self.assertTrue(any(
            "duplicate visual asset" in reason
            for reason in result["quality"]["rejection_reasons"]
        ), result)

    def test_slides_agent_builds_and_verifies_evidence_backed_deck(self):
        from PIL import Image

        image_a = self.root / "robot-a.png"
        image_b = self.root / "robot-b.png"
        Image.effect_noise((960, 540), 58).convert("RGB").save(image_a)
        Image.effect_noise((960, 540), 82).convert("RGB").save(image_b)

        sections = []
        for index in range(5):
            sections.extend([
                f"SLIDE: Robotics automation evidence {index + 1}",
                "BODY: " + (
                    "Robotics automation combines perception, planning, control, and verification. "
                    "This evidence-backed section explains measurable deployment tradeoffs, operator "
                    "oversight, failure recovery, and safe task completion. "
                ) * 2,
                f"IMAGE: {image_a if index % 2 == 0 else image_b}",
                f"LINK: Direct robotics source | https://example.com/robotics/{index + 1}",
            ])
        spec = "\n".join([
            "TITLE: Robotics Automation Field Guide",
            "SUBTITLE: Evidence-backed systems and deployment",
            "THEME: dark cyan",
            *sections,
        ])
        result = self.actions._specialist_hub().execute({
            "agent": "slides",
            "request": "Create a robotics automation PPTX with images, hyperlinks, and web research.",
            "evidence": "WEB_RESEARCH:\nFetched pages: 3\nReadable text:\nverified robotics automation research",
            "spec": spec,
        }, "chat")
        self.assertEqual(result["status"], "verified", result)
        self.assertTrue(Path(result["path"]).exists())
        self.assertGreaterEqual(result["quality"]["slides"], 6)
        self.assertGreaterEqual(result["quality"]["pictures"], 2)
        self.assertGreaterEqual(result["quality"]["direct_source_links"], 1)

    def test_thin_slides_agent_result_requires_repair(self):
        result = self.actions._specialist_hub().execute({
            "agent": "slides",
            "request": "Create a moon mission PPTX with images and links.",
            "evidence": "WEB_RESEARCH:\nFetched pages: 1\nReadable text:\nmoon mission evidence",
            "spec": "TITLE: Moon Mission\nSLIDE: Overview\nBODY: Very short.",
        }, "chat")
        self.assertEqual(result["status"], "needs_repair")
        self.assertFalse(result["quality"]["ok"])
        self.assertTrue(
            any("fewer than 5 content slides" in reason for reason in result["quality"]["reasons"]),
            result,
        )

    def test_placeholder_slide_spec_is_rejected_before_rendering(self):
        spec = "\n".join([
            "TITLE: Naruto Characters",
            "SLIDE: Scope and Purpose",
            "BODY: Exact brief and requested scope from the current request.",
            "SLIDE: Context",
            "BODY: Available evidence will be verified later.",
            "SLIDE: Major Developments",
            "BODY: KIRA OS generated visual support.",
            "SLIDE: Key Perspectives",
            "BODY: The topic remains centered on the request.",
            "SLIDE: Synthesis",
            "BODY: External claims should be verified.",
        ])
        result = self.actions._specialist_hub().execute({
            "agent": "slides",
            "request": "Create a presentation about Naruto characters.",
            "spec": spec,
        }, "chat")
        self.assertEqual(result["status"], "needs_repair")
        self.assertIn("internal process or placeholder language", result["error"])
        self.assertFalse(list((self.root / "workspace").glob("*.pptx")))

    def test_pptx_directives_do_not_leak_into_visible_slide_text(self):
        from pptx import Presentation

        sections = []
        for index in range(5):
            sections.extend([
                f"SLIDE: Character {index + 1}",
                "BODY: A focused character profile covering motivation, relationships, conflict, and development.",
                f"BULLET: Defining trait {index + 1}",
                f"BULLET: Story contribution {index + 1}",
            ])
        spec = "\n".join([
            "TITLE: Character Field Guide",
            "SUBTITLE: Narrative roles and development",
            *sections,
        ])
        result = self.actions._specialist_hub().execute({
            "agent": "slides",
            "request": "Create a substantive character presentation.",
            "spec": spec,
        }, "chat")
        self.assertEqual(result["status"], "verified", result)
        deck = Presentation(result["path"])
        visible_text = "\n".join(
            shape.text
            for slide in deck.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        self.assertNotIn("BULLET:", visible_text)
        self.assertIn("Defining trait 1", visible_text)

    def test_production_subagent_registry_filters_stress_fixtures(self):
        registry_path = Path(self.actions._subagent_registry_path())
        registry_path.parent.mkdir(parents=True, exist_ok=True)
        fake_model = self.root / "stress_lab" / "fake_model"
        fake_model.mkdir(parents=True)
        registry_path.write_text(
            """{
  "version": "subagent-registry-1",
  "subagents": [
    {
      "name": "fake",
      "backend": "mlx",
      "model_path": "%s"
    }
  ]
}
""" % fake_model,
            encoding="utf-8",
        )
        loaded = self.actions._load_subagent_registry()
        self.assertEqual(loaded["subagents"], [])

    def test_docx_quality_gate_rejects_thin_and_accepts_substantive(self):
        from docx import Document

        thin_path = self.root / "thin.docx"
        thin = Document()
        thin.add_paragraph("Short.")
        thin.save(thin_path)
        self.assertFalse(self.actions._docx_quality_report(str(thin_path), "robotics report")["ok"])

        rich_path = self.root / "rich.docx"
        rich = Document()
        for index in range(6):
            rich.add_paragraph(
                f"Robotics automation section {index + 1}. "
                + "Evidence, system design, operator oversight, verification, and recovery are analyzed. " * 3
            )
        rich.save(rich_path)
        report = self.actions._docx_quality_report(str(rich_path), "robotics automation report")
        self.assertTrue(report["ok"], report)
        self.assertGreater(report["text_chars"], 500)

    def test_docx_content_is_kept_when_link_fields_are_present(self):
        from docx import Document

        result = self.actions._generate_native_docx(
            """TITLE: Evidence Brief
SUBTITLE: Generator regression probe
CONTENT:
# Verification
The document body must remain present when a source link follows it.
- Body evidence is retained.
- Markdown bullets become document paragraphs.
LINK: Direct source | https://example.com/evidence
"""
        )
        path = self.actions._extract_artifact_path_from_text(result)
        self.assertTrue(path, result)
        document = Document(path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Verification", text)
        self.assertIn("Body evidence is retained.", text)
        hyperlinks = [
            rel for rel in document.part.rels.values()
            if rel.reltype.endswith("/hyperlink") and rel.is_external
        ]
        self.assertEqual(len(hyperlinks), 1)

    def test_docx_quality_requires_unique_usable_requested_images(self):
        from docx import Document
        from docx.shared import Inches
        from PIL import Image

        image_path = self.root / "robotics-workcell.png"
        Image.effect_noise((960, 540), 72).convert("RGB").save(image_path)
        docx_path = self.root / "duplicate-visuals.docx"
        document = Document()
        for index in range(6):
            document.add_paragraph(
                f"Robotics automation section {index + 1}. "
                + "Evidence, controls, operator oversight, verification, and recovery are analyzed. " * 3
            )
        document.add_picture(str(image_path), width=Inches(4.0))
        document.add_picture(str(image_path), width=Inches(4.0))
        document.save(docx_path)

        report = self.actions._docx_quality_report(
            str(docx_path),
            "Create a robotics automation report with images.",
        )

        self.assertFalse(report["ok"], report)
        self.assertEqual(report["unique_usable_pictures"], 1)
        self.assertIn("document-quality visuals lack variety", report["reasons"])

    def test_docx_generator_preserves_portrait_image_aspect_ratio(self):
        from docx import Document
        from PIL import Image

        portrait = self.root / "robotics-portrait.png"
        Image.effect_noise((800, 1200), 70).convert("RGB").save(portrait)
        result = self.actions._generate_native_docx(
            f"""TITLE: Robotics Portrait Brief
SUBTITLE: Aspect ratio regression
CONTENT:
HEADING: Robotics automation
PARAGRAPH: Robotics systems combine perception, planning, control, and verification for dependable operation.
PARAGRAPH: Operators define objectives, monitor execution, and review evidence before consequential changes.
PARAGRAPH: Recovery paths preserve state, expose failures, and support measured retries during field deployment.
PARAGRAPH: Validation compares requested outcomes with observable files, logs, and system state after execution.
IMAGE: {portrait}
"""
        )
        path = self.actions._extract_artifact_path_from_text(result)
        self.assertTrue(path, result)
        document = Document(path)
        self.assertEqual(len(document.inline_shapes), 1)
        shape = document.inline_shapes[0]
        ratio = float(shape.width) / float(shape.height)
        self.assertAlmostEqual(ratio, 800 / 1200, delta=0.03)

    def test_docx_structured_directives_do_not_leak_into_body_text(self):
        from docx import Document

        result = self.actions._generate_native_docx(
            """TITLE: Robotics Operations Brief
SUBTITLE: Structured content regression
CONTENT:
HEADING: Operational context
PARAGRAPH: Robotics systems combine perception, planning, control, and verification for dependable operation.
PARAGRAPH: Operators define objectives, monitor execution, and review evidence before consequential changes.
PARAGRAPH: Recovery paths preserve state, expose failures, and support measured retries during field deployment.
PARAGRAPH: Validation compares requested outcomes with observable files, logs, and system state after execution.
BULLET: Inspect observable state before changing it.
BULLET: Verify the resulting state before reporting completion.
LINK: Direct robotics source | https://example.com/robotics
"""
        )
        path = self.actions._extract_artifact_path_from_text(result)
        self.assertTrue(path, result)
        document = Document(path)
        visible_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("PARAGRAPH:", visible_text)
        self.assertNotIn("BULLET:", visible_text)
        self.assertNotIn("LINK:", visible_text)
        self.assertEqual(visible_text.count("Robotics systems combine"), 1)

    def test_global_artifact_gate_applies_to_docx(self):
        from docx import Document

        thin_path = self.root / "thin-global.docx"
        thin = Document()
        thin.add_paragraph("Claimed complete, but empty.")
        thin.save(thin_path)

        ok, report = self.actions._artifact_meets_request(
            f"File: `{thin_path}`",
            "Create a substantive robotics report.",
        )
        self.assertFalse(ok)
        self.assertIn("insufficient substantive content", report["reasons"])


if __name__ == "__main__":
    unittest.main()
