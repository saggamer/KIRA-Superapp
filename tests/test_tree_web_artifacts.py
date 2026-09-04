import os
import queue
import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch


APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

from actions import KiraActionsMixin
from interface import KiraBrain


class HarnessActions(KiraActionsMixin):
    def __init__(self, root):
        self.app_root = str(root)
        self.agentic_workspace = str(root / "workspace")
        self.agentic_logs_path = str(root / "logs")
        self.home_path = str(root)
        self.active_chat_id = "test-chat"
        self.response_queue = queue.Queue()
        self.last_artifacts = []
        self.last_artifact_path = ""
        self.artifacts_index_path = str(root / "artifacts.json")
        self.web_evidence_by_chat = {}
        self.computer_roots = {
            "home": str(root),
            "agentic_workspace": self.agentic_workspace,
        }
        os.makedirs(self.agentic_workspace, exist_ok=True)
        os.makedirs(self.agentic_logs_path, exist_ok=True)

    def _log_agentic_event(self, *_args, **_kwargs):
        return None

    def _load_chat_messages(self, _chat_id):
        return []


class TreeWebArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.actions = HarnessActions(self.root)

    def tearDown(self):
        self.temp.cleanup()

    def test_web_query_extracts_subject_instead_of_workflow(self):
        prompt = (
            "research on the best new anime to watch and generate "
            "a document with images and Crunchyroll hyperlinks browse the web"
        )
        query = self.actions._web_query_from_prompt(prompt)
        self.assertIn("best new anime to watch", query.lower())
        self.assertIn("crunchyroll", query.lower())
        self.assertNotIn("generate a document", query.lower())

    def test_web_query_stops_before_artifact_requirements(self):
        query = self.actions._web_query_from_prompt(
            "gen a ppt on the show hunterxhunter include streaming webpages links "
            "and images of characters from the anime and open it"
        )
        self.assertIn("hunterxhunter", query.lower())
        self.assertNotIn("include streaming", query.lower())

    def test_empty_search_markers_are_not_evidence(self):
        empty = (
            "WEB_RESEARCH:\nFetched pages: 0\n"
            "WEB_SEARCH:\nReal result links:\n(no result links extracted)"
        )
        self.assertFalse(self.actions._has_web_evidence(empty))
        self.assertTrue(self.actions._has_web_evidence("WEB_RESEARCH:\nFetched pages: 2\nReadable text:\nreal content"))

    def test_search_link_parser_unwraps_redirects(self):
        html = (
            '<a href="/l/?uddg=https%3A%2F%2Fexample.com%2Fstory">Story</a>'
            '<a href="https://www.google.com/url?q=https%3A%2F%2Fexample.org%2Fpage">Page</a>'
        )
        links = self.actions._extract_web_links(html, "https://html.duckduckgo.com/html/?q=test", limit=4)
        self.assertEqual([item["url"] for item in links], ["https://example.com/story", "https://example.org/page"])

    def test_web_open_uses_system_default_browser(self):
        with patch("actions.webbrowser.open", return_value=True) as browser_open:
            result = self.actions._web_open_tool("https://example.com/reference")

        browser_open.assert_called_once_with(
            "https://example.com/reference",
            new=2,
            autoraise=True,
        )
        self.assertIn("system default browser", result)

    def test_web_browse_opens_visibly_and_fetches_with_reader(self):
        opened = []
        fetched = []
        self.actions._web_open_tool = lambda target: opened.append(target) or "WEB_OPEN: visible"
        self.actions._web_fetch_tool = (
            lambda target, max_chars="12000": fetched.append((target, max_chars)) or "WEB_FETCH: readable source"
        )

        result = self.actions._run_web_search_tools(
            "[WEB_BROWSE]\nURL: https://example.com/article\nMAX_CHARS: 9000\n[/WEB_BROWSE]"
        )

        self.assertEqual(opened, ["https://example.com/article"])
        self.assertEqual(fetched, [("https://example.com/article", "9000")])
        self.assertIn("WEB_OPEN: visible", result)
        self.assertIn("WEB_FETCH: readable source", result)

    def test_web_fetch_failure_is_reported_as_failure(self):
        self.actions._web_fetch_html_source = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("network unavailable")
        )
        result = self.actions._web_fetch_tool("https://example.com")
        self.assertIn("WEB_FETCH ERROR", result)
        self.assertIn("network unavailable", result)

    def test_web_research_does_not_count_failed_page_fetches_as_evidence(self):
        self.actions._web_search_results = lambda *_args, **_kwargs: (
            [{"title": "Result", "url": "https://example.com/article"}],
            "https://example.com/search",
        )
        self.actions._web_fetch_html_source = lambda *_args, **_kwargs: (_ for _ in ()).throw(
            RuntimeError("reader unavailable")
        )

        result = self.actions._web_research_tool("test topic", max_pages=1)

        self.assertIn("WEB_RESEARCH ERROR", result)
        self.assertIn("Fetched pages: 0", result)
        self.assertIn("reader unavailable", result)
        self.assertFalse(self.actions._has_web_evidence(result))

    def test_web_research_and_images_run_in_parallel(self):
        def delayed_research(_query, max_pages="4", visible=False):
            time.sleep(0.2)
            return "WEB_RESEARCH:\nFetched pages: 1\nReadable text:\nsource"

        def delayed_images(_query, limit="4"):
            time.sleep(0.2)
            return "WEB_IMAGE_SEARCH:\nDownloaded images: 1"

        self.actions._web_research_tool = delayed_research
        self.actions._web_image_search_tool = delayed_images
        blocks = (
            "[WEB_RESEARCH]\nQUERY: robotics\n[/WEB_RESEARCH]\n"
            "[WEB_IMAGE_SEARCH]\nQUERY: robotics\n[/WEB_IMAGE_SEARCH]"
        )
        started = time.perf_counter()
        result = self.actions._run_web_search_tools(blocks)
        elapsed = time.perf_counter() - started
        self.assertLess(elapsed, 0.34)
        self.assertIn("Fetched pages: 1", result)
        self.assertIn("Downloaded images: 1", result)

    def test_pptx_quality_rejects_thin_deck_and_accepts_substantive_deck(self):
        from pptx import Presentation
        from pptx.util import Inches

        thin_path = self.root / "thin.pptx"
        thin = Presentation()
        thin.slides.add_slide(thin.slide_layouts[6])
        thin.save(thin_path)
        self.assertFalse(self.actions._pptx_quality_report(str(thin_path), "pptx with images and hyperlinks")["ok"])

        rich_path = self.root / "rich.pptx"
        rich = Presentation()
        for index in range(6):
            slide = rich.slides.add_slide(rich.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(8), Inches(4))
            run = box.text_frame.paragraphs[0].add_run()
            run.text = (f"Topic-specific slide {index + 1}. " + "Substantive evidence and analysis. " * 8)
            run.hyperlink.address = "https://example.com/source"
            if index < 2:
                image = self.root / f"image-{index}.png"
                from PIL import Image
                Image.effect_noise((960, 540), 52 + index * 22).convert("RGB").save(image)
                slide.shapes.add_picture(str(image), Inches(9), Inches(1), width=Inches(2))
        rich.save(rich_path)
        self.assertTrue(self.actions._pptx_quality_report(str(rich_path), "pptx with images and hyperlinks")["ok"])

    def test_pdf_quality_rejects_missing_requested_visuals_and_links(self):
        from fpdf import FPDF

        pdf_path = self.root / "text-only-robotics.pdf"
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=11)
        body = (
            "Robotics automation combines perception, planning, control, operator oversight, "
            "verification, and recovery. The report compares system evidence and deployment outcomes. "
        ) * 8
        pdf.multi_cell(0, 7, body)
        pdf.output(str(pdf_path))

        report = self.actions._pdf_quality_report(
            str(pdf_path),
            "Create a robotics automation PDF with images and source hyperlinks.",
        )

        self.assertFalse(report["ok"], report)
        self.assertIn("requested visuals are missing", report["reasons"])
        self.assertIn("requested hyperlinks are missing", report["reasons"])

    def test_pdf_to_pptx_preserves_portrait_page_aspect_ratio(self):
        import fitz
        from pptx import Presentation

        pdf_path = self.root / "portrait-source.pdf"
        document = fitz.open()
        page = document.new_page(width=300, height=600)
        page.insert_text((40, 80), "Portrait source page")
        document.save(pdf_path)
        document.close()
        self.actions._record_artifact = lambda *_args, **_kwargs: None

        result = self.actions._pdf_to_pptx(str(pdf_path), "Portrait conversion", "5")
        pptx_path = self.actions._extract_artifact_path_from_text(result)

        self.assertTrue(pptx_path, result)
        deck = Presentation(pptx_path)
        self.assertEqual(len(deck.slides), 1)
        picture = next(shape for shape in deck.slides[0].shapes if shape.shape_type == 13)
        ratio = float(picture.width) / float(picture.height)
        self.assertAlmostEqual(ratio, 0.5, delta=0.02)

    def test_direct_generic_artifact_shortcut_is_off_by_default(self):
        brain = KiraBrain.__new__(KiraBrain)
        old = os.environ.pop("KIRA_ENABLE_DIRECT_ARTIFACT_FALLBACK", None)
        try:
            self.assertEqual(brain._handle_direct_artifact_request("create a pptx about space", "chat"), "")
        finally:
            if old is not None:
                os.environ["KIRA_ENABLE_DIRECT_ARTIFACT_FALLBACK"] = old

    def test_creation_prompt_with_open_at_end_is_not_swallowed_as_followup(self):
        self.actions.last_artifacts = []
        result = self.actions._handle_artifact_followup(
            "gen a ppt on Hunter x Hunter with images and open it",
            "test-chat",
        )
        self.assertEqual(result, "")

    def test_preference_based_open_request_goes_to_orchestrator(self):
        self.actions._find_app_bundle = lambda _target: "/Applications/Cursor.app"
        self.actions._open_target_tool = lambda target: f"opened {target}"
        result = self.actions._handle_natural_open_close(
            "open the best designed coding app in my laptop",
            "test-chat",
        )
        self.assertEqual(result, "")

    def test_concrete_installed_app_can_open_without_model_roundtrip(self):
        self.actions._find_app_bundle = lambda target: "/Applications/Cursor.app" if target == "Cursor" else ""
        self.actions._open_target_tool = lambda target: f"opened {target}"
        result = self.actions._handle_natural_open_close("open Cursor", "test-chat")
        self.assertEqual(result, "opened Cursor")

    def test_composite_app_research_is_not_mistaken_for_artifact_followup(self):
        self.actions.last_artifacts = []
        result = self.actions._handle_artifact_followup(
            "figure out which coding app has the best design and open it too",
            "test-chat",
        )
        self.assertEqual(result, "")

    def test_app_choice_uses_inventory_branch_without_literal_open(self):
        prompt = "figure out which coding app has the best design and open it too"
        blocks = self.actions._build_deterministic_agentic_blocks(prompt)
        self.assertIn("[APP_LIST]", blocks)
        self.assertNotIn("[OPEN]", blocks)
        self.assertTrue(self.actions._prompt_requests_app_choice(prompt))

    def test_ambiguous_open_block_returns_inventory_instead_of_opening_phrase(self):
        opened = []
        self.actions._open_target_tool = lambda target: opened.append(target) or f"Opened: {target}"
        self.actions._app_list_tool = lambda: "APP_LIST:\nCursor | /Applications/Cursor.app"
        result = self.actions._run_open_tools(
            "[OPEN]\nTARGET: the best designed coding app in my laptop\n[/OPEN]"
        )
        self.assertEqual(opened, [])
        self.assertIn("APP_TARGET_RESOLUTION_REQUIRED", result)
        self.assertIn("/Applications/Cursor.app", result)

    def test_generate_it_restores_previous_artifact_request(self):
        brain = KiraBrain.__new__(KiraBrain)
        prior = (
            "gen a ppt on the show hunterxhunter include streaming webpages links "
            "and images of characters from the anime and open it"
        )
        brain._load_chat_messages = lambda _chat_id: [
            {"role": "user", "content": prior},
            {"role": "assistant", "content": "No artifact found."},
        ]
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]
        resolved = brain._resolve_contextual_task_prompt("Generate it", "test-chat")
        self.assertIn("hunterxhunter", resolved)
        self.assertIn("streaming webpages", resolved)
        self.assertIn("Follow-up instruction: Generate it", resolved)

    def test_proceed_restores_full_artifact_brief(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain._load_chat_messages = lambda _chat_id: [
            {"role": "user", "content": "create a ppt on Naruto Shippuden"},
            {
                "role": "user",
                "content": "about Naruto and Sasuke main arcs, include images and source hyperlinks in the pptx",
            },
            {"role": "assistant", "content": "Should I proceed?"},
        ]
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]
        resolved = brain._resolve_contextual_task_prompt("proceed", "test-chat")
        self.assertIn("Naruto Shippuden", resolved)
        self.assertIn("Naruto and Sasuke main arcs", resolved)
        self.assertIn("include images", resolved)
        self.assertIn("Follow-up instruction: proceed", resolved)

    def test_use_the_web_restores_unfinished_pptx_request(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain._load_chat_messages = lambda _chat_id: [
            {
                "role": "user",
                "content": "Create a ppt on Avatar: The Last Airbender and open it",
            },
            {
                "role": "assistant",
                "content": "Should I use the web before creating the presentation?",
            },
        ]
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]

        resolved = brain._resolve_contextual_task_prompt("use the web", "test-chat")

        self.assertIn("Avatar: The Last Airbender", resolved)
        self.assertIn("Create a ppt", resolved)
        self.assertIn("Follow-up instruction: use the web", resolved)

    def test_add_images_restores_unfinished_document_request(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain._load_chat_messages = lambda _chat_id: [
            {"role": "user", "content": "create a document about local AI agents"},
        ]
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]

        resolved = brain._resolve_contextual_task_prompt("add relevant images", "test-chat")

        self.assertIn("local AI agents", resolved)
        self.assertIn("Follow-up instruction: add relevant images", resolved)

    def test_promise_after_web_evidence_hands_off_to_artifact_finalizer(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain.agentic_max_steps = 3
        brain.response_queue = queue.Queue()
        brain._log_private_thoughts = lambda *_args, **_kwargs: None
        brain._emit_agent_progress = lambda *_args, **_kwargs: None
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._record_text_execution_evidence = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]
        brain._strip_private_reasoning = lambda value: str(value or "")
        brain._strip_agentic_blocks = lambda value: str(value or "")
        brain._extract_agentic_blocks = (
            lambda value: str(value or "") if "[WEB_RESEARCH]" in str(value or "") else ""
        )
        brain._agentic_branch_labels_from_text = lambda *_args, **_kwargs: []
        brain._run_agentic_capabilities_with_watchdog = lambda *_args, **_kwargs: (
            "WEB_RESEARCH:\nFetched pages: 2\nReadable text:\n"
            "Avatar: The Last Airbender is an animated fantasy series."
        )
        brain._has_raw_agentic_result = lambda value: "WEB_RESEARCH:" in str(value or "")
        brain._raw_agentic_result_needs_synthesis = lambda *_args, **_kwargs: True
        brain._agentic_goal_satisfied = lambda *_args, **_kwargs: False
        brain._continue_orchestrator_after_tool_result = lambda *_args, **_kwargs: (
            "I will now create the presentation."
        )
        brain._split_orchestrator_response = lambda value: ("", value)
        brain._wants_artifact_generation = lambda *_args, **_kwargs: True
        brain._artifact_was_really_generated = lambda *_args, **_kwargs: False

        result = brain._run_orchestrator_agent_loop(
            "Create a ppt on Avatar: The Last Airbender and open it. Use the web.",
            "",
            "[WEB_RESEARCH]\nQUERY: Avatar The Last Airbender\n[/WEB_RESEARCH]",
            "system",
            "",
            "test-chat",
            force_pathway=True,
        )

        self.assertIn("Fetched pages: 2", result)
        self.assertNotIn("I will now create", result)

    def test_create_the_ppt_now_restores_unfinished_pptx_request(self):
        brain = KiraBrain.__new__(KiraBrain)
        brain._load_chat_messages = lambda _chat_id: [
            {
                "role": "user",
                "content": "Create a ppt on Avatar: The Last Airbender and open it",
            },
            {"role": "user", "content": "use the web"},
            {"role": "assistant", "content": "I gathered web evidence."},
        ]
        brain._log_agentic_event = lambda *_args, **_kwargs: None
        brain._truncate = lambda value, limit: str(value)[:limit]

        resolved = brain._resolve_contextual_task_prompt("create the ppt now", "test-chat")

        self.assertIn("Avatar: The Last Airbender", resolved)
        self.assertIn("Follow-up instruction: create the ppt now", resolved)

    def test_artifact_research_requests_visible_browser(self):
        blocks = self.actions._build_web_evidence_blocks(
            "create a pptx about robotics with images"
        )
        self.assertIn("VISIBLE: true", blocks)

    def test_tree_web_leaves_default_to_visible_browser(self):
        artifact_blocks = self.actions._build_deterministic_agentic_blocks(
            "research robotics and create a pptx with images"
        )
        search_blocks = self.actions._build_deterministic_agentic_blocks(
            "search the web for orbital robotics"
        )

        self.assertIn("[WEB_RESEARCH]", artifact_blocks)
        self.assertIn("VISIBLE: true", artifact_blocks)
        self.assertIn("[WEB_SEARCH]", search_blocks)
        self.assertIn("VISIBLE: true", search_blocks)

    def test_plain_web_search_opens_default_browser(self):
        opened = []
        self.actions._web_open_tool = (
            lambda target: opened.append(target) or "WEB_OPEN: visible"
        )
        self.actions._web_search_results = (
            lambda *_args, **_kwargs: ([], "https://example.com/search")
        )

        self.actions._web_search_tool("orbital robotics")

        self.assertEqual(len(opened), 1)
        self.assertIn("google.com/search", opened[0])

    def test_visible_web_research_opens_default_browser_search(self):
        opened = []
        self.actions._web_open_tool = (
            lambda target: opened.append(target) or "WEB_OPEN: visible"
        )
        self.actions._web_search_results = (
            lambda *_args, **_kwargs: ([], "https://example.com/search")
        )

        self.actions._web_research_tool("robotics", max_pages=1, visible=True)

        self.assertEqual(len(opened), 1)
        self.assertIn("google.com/search", opened[0])

    def test_source_records_exclude_search_pages(self):
        brain = KiraBrain.__new__(KiraBrain)
        evidence = (
            "Search URL: https://html.duckduckgo.com/html/?q=avatar\n"
            "URL: https://en.wikipedia.org/wiki/Avatar:_The_Last_Airbender\n"
            "URL: https://www.imdb.com/title/tt0417299\n"
        )

        records = brain._source_records_from_evidence(evidence)
        urls = [item["url"] for item in records]

        self.assertEqual(len(urls), 2)
        self.assertFalse(any("duckduckgo" in url for url in urls))

    def test_artifact_fallback_refuses_unverified_planning_prose(self):
        prompt = "create a pptx about Naruto and Sasuke main arcs with images"
        evidence = (
            "WEB_RESEARCH:\nFetched pages: 2\nReadable text:\n"
            "Naruto and Sasuke follow parallel character arcs shaped by rivalry, separation, "
            "conflict, reconciliation, and their final confrontation. Their choices repeatedly "
            "change the direction of the wider story and provide the presentation's central comparison.\n"
            "https://naruto.fandom.com/wiki/Plot_of_Naruto\n"
            "https://en.wikipedia.org/wiki/Naruto\n"
        )
        self.actions._generate_prompt_specific_artifact_with_orchestrator = (
            lambda *_args, **_kwargs: "I am generating the PPTX now. Please wait while I compile the file."
        )
        self.actions._wants_artifact_generation = lambda *_args, **_kwargs: True
        result = self.actions._ensure_requested_artifact_or_real_open(
            prompt,
            "I am generating the PPTX now. Please wait while I compile the file.",
            evidence,
            "test-chat",
        )
        self.assertFalse(self.actions._artifact_was_really_generated(result))
        self.assertNotIn("I am generating", result)
        self.assertIn("did not create", result.lower())

    def test_image_search_failure_does_not_publish_placeholder_visuals(self):
        prompt = "create a pptx about Naruto and Sasuke main arcs with images"
        self.actions._build_web_evidence_blocks = lambda *_args, **_kwargs: (
            "[WEB_RESEARCH]\nQUERY: Naruto and Sasuke arcs\n[/WEB_RESEARCH]\n"
            "[WEB_IMAGE_SEARCH]\nQUERY: Naruto and Sasuke arcs\n[/WEB_IMAGE_SEARCH]"
        )
        self.actions._run_agentic_capabilities = lambda *_args, **_kwargs: (
            "WEB_RESEARCH:\nFetched pages: 0\nReadable text:\n(empty)\n"
            "WEB_IMAGE_SEARCH:\nDownloaded images: 0"
        )
        self.actions._generate_prompt_specific_artifact_with_orchestrator = (
            lambda *_args, **_kwargs: "I am generating the PPTX now."
        )
        self.actions._wants_artifact_generation = lambda *_args, **_kwargs: True
        result = self.actions._ensure_requested_artifact_or_real_open(
            prompt,
            "I am generating the PPTX now.",
            "",
            "test-chat",
        )
        self.assertFalse(self.actions._artifact_was_really_generated(result))
        self.assertNotIn("I am generating", result)
        self.assertIn("did not create", result.lower())

    def test_generic_kira_deck_fails_topic_and_source_gate(self):
        deck = APP_DIR / "kira_agentic_workspace" / "KIRA_Deck.pptx"
        if not deck.exists():
            self.skipTest("Observed regression artifact is unavailable")
        report = self.actions._pptx_quality_report(
            str(deck),
            "gen a ppt on hunterxhunter with character images and streaming website hyperlinks",
        )
        self.assertFalse(report["ok"])
        self.assertIn("links only point to generic search pages, not direct sources", report["reasons"])
        self.assertTrue(
            any(
                "generic" in reason or "placeholder" in reason or "internal process" in reason
                for reason in report["reasons"]
            ),
            report["reasons"],
        )


if __name__ == "__main__":
    unittest.main()
